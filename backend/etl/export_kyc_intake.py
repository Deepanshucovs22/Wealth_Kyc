"""
Export every row of the New Account tables (schema kyc_intake) to Excel.

    python backend/etl/export_kyc_intake.py            -> exports/kyc_intake_<timestamp>.xlsx

One sheet per table, every column. Uploaded files are not copied into the
workbook — `content` shows the byte count, and `sha256` identifies the file.
The workbook holds personal data: exports/ is git-ignored, keep it that way.
"""
from __future__ import annotations

import json
import sys
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import psycopg2
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from backend.config import app_dsn, describe_target  # noqa: E402

TABLES = [
    ("kyc_session", "SELECT * FROM kyc_intake.kyc_session ORDER BY created_at"),
    ("kyc_form_data", "SELECT * FROM kyc_intake.kyc_form_data ORDER BY submitted_at"),
    ("kyc_document", """SELECT document_id, session_id, doc_type, seq, file_name, mime_type,
                               size_bytes, sha256, octet_length(content) AS content_bytes, uploaded_at
                        FROM kyc_intake.kyc_document ORDER BY session_id, document_id"""),
    ("kyc_ocr_result", "SELECT * FROM kyc_intake.kyc_ocr_result ORDER BY session_id, doc_type, attempt"),
    ("kyc_session_event", "SELECT * FROM kyc_intake.kyc_session_event ORDER BY session_id, at, event_id"),
]
CELL_LIMIT = 32_000          # Excel refuses more than 32,767 characters in a cell


def cell(v):
    if isinstance(v, datetime):
        return v.astimezone().replace(tzinfo=None) if v.tzinfo else v
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, (dict, list)):
        v = json.dumps(v, ensure_ascii=False)
    if isinstance(v, (bytes, memoryview)):
        return f"<{len(bytes(v))} bytes>"
    if isinstance(v, str) and len(v) > CELL_LIMIT:
        return v[:CELL_LIMIT] + f" … [truncated, {len(v):,} characters in the database]"
    return v


def main() -> None:
    out_dir = Path(__file__).resolve().parents[2] / "exports"
    out_dir.mkdir(exist_ok=True)
    out = out_dir / f"kyc_intake_{datetime.now():%Y%m%d_%H%M%S}.xlsx"

    wb = Workbook()
    wb.remove(wb.active)
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2F6FEB")

    with psycopg2.connect(app_dsn()) as conn, conn.cursor() as cur:
        for name, sql in TABLES:
            cur.execute(sql)
            cols = [d[0] for d in cur.description]
            data = cur.fetchall()
            ws = wb.create_sheet(name)
            ws.append(cols)
            for c in ws[1]:
                c.font, c.fill = head_font, head_fill
            for row in data:
                ws.append([cell(v) for v in row])
            for i, col in enumerate(cols, 1):
                width = max([len(col)] + [min(len(str(r[i - 1] or "")), 60) for r in data])
                ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 62)
            for r in ws.iter_rows(min_row=2):
                for c in r:
                    c.alignment = Alignment(vertical="top", wrap_text=isinstance(c.value, str) and len(c.value) > 60)
            ws.freeze_panes = "B2"
            print(f"  {name:20} {len(data):4} rows  {len(cols):3} columns")

    wb.save(out)
    print(f"\n  {describe_target()} -> {out}")


if __name__ == "__main__":
    main()
